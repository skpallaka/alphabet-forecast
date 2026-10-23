"""Day 2 + 5: build the Excel P&L model with live formulas.

Sheets
  Assumptions   fade factor, scenario growth deltas (the only inputs besides history)
  Model_Base / Model_Bull / Model_Bear   identical layouts: history (blue = hard-coded from SEC) + forecast formulas
  Scenarios     FY26E / FY27E side by side, all formulas into the three model sheets
  Seasonality   each quarter's share of annual revenue (why the YoY method needs same-quarter bases)

Forecast logic (all in-cell formulas, nothing pasted):
  segment revenue(t) = segment revenue(t-4) x (1 + growth(t))
  growth(t)          = latest actual YoY x fade^k  + scenario delta          (k = quarters ahead)
  cost of revenues, S&M = % of revenue, trailing 4-quarter average of the ratio
  R&D, G&A           = % of revenue in Base; held at Base-case dollars in Bull/Bear (largely fixed commitments)
"""
import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

ROOT = Path(__file__).resolve().parent.parent
FIRST = pd.Timestamp("2019-03-31")
N_HIST = 30                  # Q1'19 .. Q2'26
N_FC = 6                     # Q3'26 .. Q4'27
C0 = 3                       # first quarter column (C)
LAST_ACT = C0 + N_HIST - 1   # AF
LAST = LAST_ACT + N_FC       # AL

BLUE, BLACK, GREEN = Font(color="0000FF"), Font(color="000000"), Font(color="007A33")
BOLD = Font(bold=True)
HDR_FILL = PatternFill("solid", fgColor="1F3864")
FC_FILL = PatternFill("solid", fgColor="FFF2CC")
INPUT_FILL = PatternFill("solid", fgColor="DDEBF7")
TOP = Border(top=Side(style="thin"))

SEGMENTS = [  # label, db column, row on model sheet
    ("Search & other", "rev_search"),
    ("YouTube ads", "rev_youtube"),
    ("Network", "rev_network"),
    ("Subscriptions, platforms & devices", "rev_subscriptions"),
    ("Google Cloud", "rev_cloud"),
    ("Other Bets", "rev_other_bets"),
]
# row map on model sheets
R = dict(hdr=3, date=4, flag=5, k=6,
         seg0=8, hedge=14, total=15,
         yoy0=18, yoy_total=24,
         cor=27, gp=28, gm=29, rd=30, sm=31, ga=32, ec=33, opex=34, oi=35, om=36,
         rat_cor=39, rat_rd=40, rat_sm=41, rat_ga=42,
         gro0=45)
# assumption sheet: scenario deltas table
A_FADE = "Assumptions!$C$4"
A_DELTA_ROW0 = 9     # rows 9..14 = segments; cols C,D,E = Bear, Base, Bull
SCEN_COL = {"Bear": "C", "Base": "D", "Bull": "E"}
DELTAS = {  # percentage points added to each forecast quarter's YoY growth
    "rev_search": (-0.02, 0, 0.02), "rev_youtube": (-0.02, 0, 0.02), "rev_network": (0, 0, 0),
    "rev_subscriptions": (0, 0, 0), "rev_cloud": (-0.05, 0, 0.05), "rev_other_bets": (0, 0, 0),
}


def qlabel(ts):
    return f"Q{(ts.month - 1) // 3 + 1}'{str(ts.year)[2:]}"


def build_model(ws, df, scen, base_title=None):
    ws.sheet_view.showGridLines = False
    ws["A1"] = f"Alphabet Inc. - quarterly P&L model ({scen} case)  |  $ millions"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = ("Blue = hard-coded SEC history. Black = formula. Yellow-shaded columns = forecast. "
                "Source: SEC EDGAR XBRL (companyfacts + 10-Q/10-K instance documents).")
    ws["A2"].font = Font(italic=True, color="595959")
    quarters = pd.date_range(FIRST, periods=N_HIST + N_FC, freq="QE")
    for i, q in enumerate(quarters):
        c = C0 + i
        fc = i >= N_HIST
        ws.cell(R["hdr"], c, qlabel(q))
        ws.cell(R["date"], c, q.to_pydatetime()).number_format = "mmm-yy"
        ws.cell(R["flag"], c, "Forecast" if fc else "Actual")
        ws.cell(R["k"], c, i - N_HIST + 1 if fc else None)
        for r in (R["hdr"], R["date"], R["flag"], R["k"]):
            cell = ws.cell(r, c)
            cell.alignment = Alignment(horizontal="right")
            if r == R["hdr"]:
                cell.font = Font(bold=True, color="FFFFFF"); cell.fill = HDR_FILL
            elif fc:
                cell.fill = FC_FILL
    ws.cell(R["hdr"], 1, "Quarter").font = Font(bold=True, color="FFFFFF"); ws.cell(R["hdr"], 1).fill = HDR_FILL
    ws.cell(R["hdr"], 2).fill = HDR_FILL
    ws.cell(R["k"], 1, "Quarters ahead (k)").font = Font(color="7F7F7F")

    def label(r, text, bold=False, indent=0):
        c = ws.cell(r, 1, text)
        c.font = Font(bold=bold)
        c.alignment = Alignment(indent=indent)

    def put(r, i, value, fmt="#,##0;(#,##0)", hist_input=False, bold=False):
        c = ws.cell(r, C0 + i, value)
        c.number_format = fmt
        c.font = Font(bold=bold, color="0000FF" if hist_input else "000000")
        if i >= N_HIST:
            c.fill = FC_FILL
        return c

    pct = "0.0%;(0.0%)"
    is_base = scen == "Base"
    label(R["seg0"] - 1, "REVENUE", bold=True)
    for j, (name, col) in enumerate(SEGMENTS):
        r = R["seg0"] + j
        label(r, name, indent=1)
        for i in range(N_HIST + N_FC):
            cl = L(C0 + i)
            if i < N_HIST:
                v = df[col].iloc[i]
                put(r, i, round(float(v), 1) if is_base else f"=Model_Base!{cl}{r}", hist_input=is_base)
                if not is_base:
                    ws.cell(r, C0 + i).font = GREEN
            else:
                g = f"{cl}{R['gro0'] + j}"
                put(r, i, f"={L(C0 + i - 4)}{r}*(1+{g})")
    label(R["hedge"], "Hedging gains / (losses)", indent=1)
    label(R["total"], "Total revenues", bold=True)
    for i in range(N_HIST + N_FC):
        cl = L(C0 + i)
        if i < N_HIST:
            put(R["hedge"], i, round(float(df["rev_hedging"].iloc[i]), 1) if is_base else f"=Model_Base!{cl}{R['hedge']}",
                hist_input=is_base)
            if not is_base:
                ws.cell(R["hedge"], C0 + i).font = GREEN
        else:
            put(R["hedge"], i, 0)          # not forecast
        c = put(R["total"], i, f"=SUM({cl}{R['seg0']}:{cl}{R['hedge']})", bold=True)
        c.border = TOP

    label(R["yoy0"] - 1, "YoY GROWTH", bold=True)
    for j, (name, _) in enumerate(SEGMENTS):
        r = R["yoy0"] + j
        label(r, name, indent=1)
        for i in range(4, N_HIST + N_FC):
            put(r, i, f"={L(C0 + i)}{R['seg0'] + j}/{L(C0 + i - 4)}{R['seg0'] + j}-1", pct)
    label(R["yoy_total"], "Total revenues", bold=True)
    for i in range(4, N_HIST + N_FC):
        put(R["yoy_total"], i, f"={L(C0 + i)}{R['total']}/{L(C0 + i - 4)}{R['total']}-1", pct, bold=True)

    label(R["cor"] - 1, "PROFIT & LOSS", bold=True)
    names = {"cor": ("Cost of revenues", "cost_of_revenues"), "rd": ("Research & development", "rd"),
             "sm": ("Sales & marketing", "sales_marketing"), "ga": ("General & administrative", "g_and_a"),
             "ec": ("European Commission fines", "ec_fines")}
    for key, (text, col) in names.items():
        label(R[key], text, indent=1)
    label(R["gp"], "Gross profit", bold=True)
    label(R["gm"], "  Gross margin")
    label(R["opex"], "Total operating expenses (R&D + S&M + G&A + fines)", bold=True)
    label(R["oi"], "Operating income", bold=True)
    label(R["om"], "  Operating margin")
    ratio_row = {"cor": R["rat_cor"], "rd": R["rat_rd"], "sm": R["rat_sm"], "ga": R["rat_ga"]}
    for i in range(N_HIST + N_FC):
        cl = L(C0 + i)
        for key, (text, col) in names.items():
            r = R[key]
            if i < N_HIST:
                v = round(float(df[col].iloc[i]), 1)
                put(r, i, v if is_base else f"=Model_Base!{cl}{r}", hist_input=is_base)
                if not is_base:
                    ws.cell(r, C0 + i).font = GREEN
            elif key == "ec":
                put(r, i, 0)
            elif not is_base and key in ("rd", "ga"):
                put(r, i, f"=Model_Base!{cl}{r}")           # fixed at Base dollars
                ws.cell(r, C0 + i).font = GREEN
            else:
                put(r, i, f"={cl}{R['total']}*{cl}{ratio_row[key]}")
        c = put(R["gp"], i, f"={cl}{R['total']}-{cl}{R['cor']}", bold=True); c.border = TOP
        put(R["gm"], i, f"={cl}{R['gp']}/{cl}{R['total']}", pct)
        put(R["opex"], i, f"={cl}{R['rd']}+{cl}{R['sm']}+{cl}{R['ga']}+{cl}{R['ec']}", bold=True)
        c = put(R["oi"], i, f"={cl}{R['gp']}-{cl}{R['opex']}", bold=True); c.border = TOP
        put(R["om"], i, f"={cl}{R['oi']}/{cl}{R['total']}", pct)

    label(R["rat_cor"] - 1, "COST RATIOS (% of revenue; forecast = trailing 4-quarter average)", bold=True)
    for key, text in (("cor", "Cost of revenues"), ("rd", "R&D"), ("sm", "Sales & marketing"), ("ga", "G&A")):
        r = ratio_row[key]
        label(r, text, indent=1)
        for i in range(N_HIST + N_FC):
            cl = L(C0 + i)
            if i < N_HIST:
                put(r, i, f"={cl}{R[key]}/{cl}{R['total']}", pct)
            else:
                put(r, i, f"=AVERAGE({L(C0 + i - 4)}{r}:{L(C0 + i - 1)}{r})", pct)

    label(R["gro0"] - 1, "GROWTH ASSUMPTIONS = latest actual YoY x fade^k + scenario delta", bold=True)
    for j, (name, col) in enumerate(SEGMENTS):
        r = R["gro0"] + j
        label(r, name, indent=1)
        for i in range(N_HIST, N_HIST + N_FC):
            cl = L(C0 + i)
            delta = f"Assumptions!${SCEN_COL[scen]}${A_DELTA_ROW0 + j}"
            put(r, i, f"=${L(LAST_ACT)}{R['yoy0'] + j}*{A_FADE}^{cl}${R['k']}+{delta}", pct)

    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 2
    for i in range(N_HIST + N_FC):
        ws.column_dimensions[L(C0 + i)].width = 10.5
    ws.freeze_panes = ws.cell(R["flag"] + 1, C0)


def build_assumptions(ws):
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Assumptions & scenario inputs"; ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Blue cells on blue fill are the only inputs. Everything else in the workbook is formula-driven."
    ws["A2"].font = Font(italic=True, color="595959")
    ws["A4"] = "Growth fade factor (per quarter ahead)"
    ws["C4"] = 0.90
    ws["C4"].font = BLUE; ws["C4"].fill = INPUT_FILL; ws["C4"].number_format = "0.00"
    ws["D4"] = ("Latest YoY growth is multiplied by fade^k for the k-th forecast quarter. 1.00 = carry growth forward unchanged. "
                "Backtest note: fading cost accuracy on Cloud (it kept accelerating) but is neutral elsewhere.")
    ws["D4"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[4].height = 48
    ws["A7"] = "Scenario deltas: percentage points added to each forecast quarter's YoY growth"
    ws["A7"].font = BOLD
    for c, t in zip("CDE", ("Bear", "Base", "Bull")):
        ws[f"{c}8"] = t; ws[f"{c}8"].font = Font(bold=True, color="FFFFFF"); ws[f"{c}8"].fill = HDR_FILL
        ws[f"{c}8"].alignment = Alignment(horizontal="center")
    ws["A8"] = "Segment"; ws["A8"].font = Font(bold=True, color="FFFFFF"); ws["A8"].fill = HDR_FILL; ws["B8"].fill = HDR_FILL
    for j, (name, col) in enumerate(SEGMENTS):
        r = A_DELTA_ROW0 + j
        ws.cell(r, 1, name)
        for c, v in zip("CDE", DELTAS[col]):
            cell = ws[f"{c}{r}"]
            cell.value = v; cell.font = BLUE; cell.fill = INPUT_FILL; cell.number_format = "+0.0%;-0.0%;0.0%"
            cell.alignment = Alignment(horizontal="center")
    notes = [
        "Other fixed assumptions (documented, not inputs):",
        "  - Cost of revenues, sales & marketing: % of revenue = trailing 4-quarter average ratio (rolling).",
        "  - R&D and G&A: same method in Base; held at Base-case dollars in Bull/Bear (headcount/compute commitments don't flex with a revenue swing).",
        "  - Hedging gains/losses and EC fines: forecast at zero.",
        "  - Seasonality: each forecast quarter grows off the same quarter a year earlier, so the Q4 ad peak is carried automatically.",
    ]
    for i, t in enumerate(notes):
        ws.cell(17 + i, 1, t).font = BOLD if i == 0 else Font(color="404040")
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 2
    for c in "CDE":
        ws.column_dimensions[c].width = 12
    ws.column_dimensions["D"].width = 12


def fy_sum(sheet, row, fy):
    """FY sum of a model-sheet row, as a formula over the 4 quarter columns of fiscal year fy."""
    i0 = (fy - 2019) * 4
    return f"SUM({sheet}!{L(C0 + i0)}{row}:{L(C0 + i0 + 3)}{row})"


def build_scenarios(ws):
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Scenario summary: Bull / Base / Bear  |  $ millions"; ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = ("Bear/Bull = Base growth -/+ the deltas on the Assumptions sheet (Cloud +/-5pp, Search and YouTube +/-2pp). "
                "FY2026E = 2 actual quarters + 2 forecast quarters.")
    ws["A2"].font = Font(italic=True, color="595959")
    heads = ["", "", "FY2025A", "FY26E Bear", "FY26E Base", "FY26E Bull", "FY27E Bear", "FY27E Base", "FY27E Bull"]
    for c, t in enumerate(heads, 1):
        cell = ws.cell(4, c, t); cell.font = Font(bold=True, color="FFFFFF"); cell.fill = HDR_FILL
        cell.alignment = Alignment(horizontal="center")
    lines = [("Google Cloud revenue", R["seg0"] + 4, "#,##0"), ("Search & other revenue", R["seg0"], "#,##0"),
             ("Total revenues", R["total"], "#,##0"), ("Gross profit", R["gp"], "#,##0"),
             ("Operating expenses", R["opex"], "#,##0"), ("Operating income", R["oi"], "#,##0")]
    row_of = {}
    for k, (text, mrow, fmt) in enumerate(lines):
        r = 5 + k
        row_of[text] = r
        ws.cell(r, 1, text).font = BOLD if text in ("Total revenues", "Operating income") else Font()
        ws.cell(r, 3, "=" + fy_sum("Model_Base", mrow, 2025)).number_format = fmt
        for c, (fy, sc) in enumerate([(2026, "Bear"), (2026, "Base"), (2026, "Bull"),
                                      (2027, "Bear"), (2027, "Base"), (2027, "Bull")], 4):
            ws.cell(r, c, "=" + fy_sum(f"Model_{sc}", mrow, fy)).number_format = fmt
    r = 11
    ws.cell(r, 1, "Operating margin").font = BOLD
    for c in range(3, 10):
        cl = L(c)
        ws.cell(r, c, f"={cl}{row_of['Operating income']}/{cl}{row_of['Total revenues']}").number_format = "0.0%"
    ws.cell(12, 1, "Total revenue growth YoY")
    ws.cell(12, 3, f"=C{row_of['Total revenues']}/{fy_sum('Model_Base', R['total'], 2024)}-1").number_format = "0.0%"
    for c in (4, 5, 6):
        ws.cell(12, c, f"={L(c)}{row_of['Total revenues']}/$C${row_of['Total revenues']}-1").number_format = "0.0%"
    for c, base_c in ((7, 4), (8, 5), (9, 6)):
        ws.cell(12, c, f"={L(c)}{row_of['Total revenues']}/{L(base_c)}{row_of['Total revenues']}-1").number_format = "0.0%"
    ws.cell(14, 1, "vs. Base case").font = BOLD
    ws.cell(15, 1, "Operating income vs Base ($M)")
    ws.cell(16, 1, "Operating income vs Base (%)")
    ws.cell(17, 1, "Total revenue vs Base ($M)")
    for c, base_c in ((4, 5), (5, 5), (6, 5), (7, 8), (8, 8), (9, 8)):
        cl, bl = L(c), L(base_c)
        ws.cell(15, c, f"={cl}{row_of['Operating income']}-{bl}{row_of['Operating income']}").number_format = "#,##0;(#,##0)"
        ws.cell(16, c, f"={cl}{row_of['Operating income']}/{bl}{row_of['Operating income']}-1").number_format = "0.0%;(0.0%)"
        ws.cell(17, c, f"={cl}{row_of['Total revenues']}-{bl}{row_of['Total revenues']}").number_format = "#,##0;(#,##0)"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 2
    for c in range(3, 10):
        ws.column_dimensions[L(c)].width = 13


def build_seasonality(ws):
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Seasonality: each quarter's share of full-year total revenue"; ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Q4 is consistently the largest quarter (holiday ad spend + Pixel launches), which is why every forecast quarter grows off the same quarter last year."
    ws["A2"].font = Font(italic=True, color="595959")
    for c, t in enumerate(["Fiscal year", "Q1", "Q2", "Q3", "Q4"], 1):
        cell = ws.cell(4, c, t); cell.font = Font(bold=True, color="FFFFFF"); cell.fill = HDR_FILL
    for k, fy in enumerate(range(2019, 2026)):
        r = 5 + k
        ws.cell(r, 1, fy)
        i0 = (fy - 2019) * 4
        for q in range(4):
            ws.cell(r, 2 + q, f"=Model_Base!{L(C0 + i0 + q)}{R['total']}/{fy_sum('Model_Base', R['total'], fy)}").number_format = "0.0%"
    ws.cell(12, 1, "Average").font = BOLD
    for q in range(4):
        cl = L(2 + q)
        c = ws.cell(12, 2 + q, f"=AVERAGE({cl}5:{cl}11)"); c.number_format = "0.0%"; c.font = BOLD; c.border = TOP
    ws.column_dimensions["A"].width = 14


def main():
    con = sqlite3.connect(ROOT / "data" / "alphabet.db")
    df = pd.read_sql("select * from v_pnl where quarter_end >= '2019-03-31'", con, parse_dates=["quarter_end"])
    assert len(df) == N_HIST and df.rev_cloud.notna().all(), "expected 30 complete quarters"
    wb = Workbook()
    ws_a = wb.active; ws_a.title = "Assumptions"
    build_assumptions(ws_a)
    for scen in ("Base", "Bull", "Bear"):
        build_model(wb.create_sheet(f"Model_{scen}"), df, scen)
    build_scenarios(wb.create_sheet("Scenarios"))
    build_seasonality(wb.create_sheet("Seasonality"))
    wb._sheets = [wb["Scenarios"], wb["Assumptions"], wb["Model_Base"], wb["Model_Bull"], wb["Model_Bear"], wb["Seasonality"]]
    wb.active = 0
    out = ROOT / "excel" / "alphabet_pnl_model.xlsx"
    out.parent.mkdir(exist_ok=True)
    wb.save(out)
    print("saved", out)


if __name__ == "__main__":
    main()
